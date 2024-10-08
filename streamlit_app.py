import datetime
import pandas as pd
import streamlit as st
import openpyxl


df = openpyxl.load_workbook('namelist.xlsx')
sheet = df.active #Assuming the data is in the first sheet

# Define CRUD functions

def add_record(id, country, company, name, designation, dietary, contact, address, vehicle, posting_date, status, golf, golf_handicap, depost_date):
    max_row = sheet.max_row + 1
    sheet['A' + str(max_row)] = id 
    sheet['B' + str(max_row)] = country
    sheet['C' + str(max_row)] = company
    sheet['D' + str(max_row)] = name
    sheet['E' + str(max_row)] = designation
    sheet['F' + str(max_row)] = dietary
    sheet['G' + str(max_row)] = contact
    sheet['H' + str(max_row)] = address
    sheet['I' + str(max_row)] = vehicle
    sheet['J' + str(max_row)] = posting_date
    sheet['K' + str(max_row)] = status
    sheet['L' + str(max_row)] = golf
    sheet['M' + str(max_row)] = golf_handicap
    df.save('namelist.xlsx')

def del_record(id):
    row_to_delete = None
    for row in df.iter_rows(values_only=True):
        if row[0] == id:
            row_to_delete = row
            break
    if row_to_delete:
        sheet.delete_rows(row_to_delete[0],1)
        df.save('namelist.xlsx')
    else:
        st.warning('Not found.')

def update_record(id, country, company, name, designation, dietary, contact, address, vehicle, posting_date, status, golf, golf_handicap, depost_date):
    for row in sheet.iter_rows(values_only=True):
        if row[0] == id:
            row[1] = country
            row[2] = company
            row[3] = name
            row[4] = designation
            row[5] = dietary
            row[6] = contact
            row[7] = address
            row[8] = vehicle
            row[9] = posting_date
            row[10]= status
            row[11]= golf
            row[12]= golf_handicap
            df.save('namelist.xlsx')
            break
    else:
        st.warning("Not found.")

def view_all_employees():
    data = []
    for row in sheet.iter_rows(values_only=True):
        data.append(row)
    return data

# Show app title and description.
st.set_page_config(
    page_title="Namelist", 
    page_icon="👤",
    layout= "wide",
)
st.title("👤 Namelist")

# Add Employee
st.header("Add New Employee")
country = st.text_input("Country")
company = st.text_input("Company")
name = st.text_input("Name")
designation = st.text_input("Designation")
dietary = st.text_input("Dietary Restriction")
contact = st.text_input("Contact No.")
address = st.text_input("Address")
vehicle = st.text_input("Vehicle No.")
posting_date = st.date_input("Posting Date")
status = st.radio("Status", ("Active", "De-posted"))
if status == "De-posted":
    depost_date = st.date_input("De-posted Date")
golf = st.radio("Plays golf?", ("Yes", "No"))
golf_handicap = st.radio("Golf Hanficap?", ("No", "Yes"))
if golf_handicap == "Yes":
    golf_handicap = st.text_input("What's the handicap?")
if st.button("Add"):
    add_record(country, company, name, designation, dietary, contact, address, vehicle, posting_date, status, golf, golf_handicap, depost_date)
    st.success("Added successfully!")

# Delete Employee
st.header("Delete Employee")
id = st.number_input("Employee ID")

if st.button("Delete"):
    del_record(id)
    st.success("Deleted successfully!")

# Update Employee
st.header("Update Employee")
#id = st.number_input("Employee ID")
country = st.text_input("New Country")
company = st.text_input("New Company")
name = st.text_input("New Name")
designation = st.text_input("New Designation")
dietary = st.text_input("New Dietary Restriction")
contact = st.text_input("New Contact No.")
address = st.text_input("New Address")
vehicle = st.text_input("New Vehicle No.")
posting_date = st.date_input("New Posting Date")
status = st.radio("New Status", ("Active", "De-posted"))
if status == "De-posted":
    depost_date = st.date_input("De-posted Date")
golf = st.radio("New Plays golf?", ("Yes", "No"))
golf_handicap = st.radio("New Golf Hanficap?", ("No", "Yes"))
if golf_handicap == "Yes":
    golf_handicap = st.text_input("What's the handicap?")

if st.button("Update"):
    update_record(id, country, company, name, designation, dietary, contact, address, vehicle, posting_date, status, golf, golf_handicap, depost_date)
    st.success("Updated successfully!")

# View All Employees
st.header("View All Employees")
employees = view_all_employees()
if employees:
    df = pd.DataFrame(employees, columns=["ID", "Country", "Company", "Name", "Designation", "Dietary Restriction", "Contact No.", "Address", "Vehicle No.", "Posting Date", "Status", "Golf", "Golf Handicap", "De-posted Date"])
    st.dataframe(df, index=False)
else:
    st.warning("No employees found.")